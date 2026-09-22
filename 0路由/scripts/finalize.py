"""P6/P7 后处理与产出：DOI 命名规范化、有效性校验、元信息表、成功失败日志。

用法：
    python finalize.py
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import _common as C  # noqa: E402

MIN_PDF_BYTES = 10_000     # 小于这个大小的"PDF"基本都是错误页
PDF_MAGIC = b"%PDF"

REASON_CN = {
    "not_found": "各源都没有可下载副本",
    "paywalled_no_access": "有付费墙且当前无可用机构通道",
    "ip_blocked": "出版商封禁出口 IP",
    "cf_challenge": "被 Cloudflare/CAPTCHA 拦截",
    "network": "网络层失败（超时/DNS/TLS/代理）",
    "bad_identifier": "DOI/arXiv 号无效或无法解析",
    "file_missing": "声称成功但磁盘上无有效 PDF",
    "duplicate": "与队列中已有条目重复",
    "unsupported": "输入类型或来源当前不支持",
    "unknown": "未归类",
}

RETRY_HINT = {
    "not_found": "换来源或开灰色源（policy.py set max）；确认 DOI 是否最新。",
    "paywalled_no_access": "配置机构通道（Elsevier API Key / WebVPN），或改用 OA 版本。",
    "ip_blocked": "降低 batch_workers、拉大 request_delay，或换出口 IP；必要时换时段重试。",
    "cf_challenge": "安装隐身浏览器后端（bootstrap.py apply --browser），或改用出版商 API 通道。",
    "network": "检查代理与 DNS；doctor.py 看端点连通性。",
    "bad_identifier": "人工核对 DOI；可用 search.py 按标题找回正确 DOI。",
    "file_missing": "直接重跑 download.py --retry-only。",
    "duplicate": "无需处理。",
    "unsupported": "人工处理。",
    "unknown": "看 failures.md 里的原始错误信息。",
}

COLUMNS = [
    ("seq", "序号"),
    ("doi", "DOI/标识符"),
    ("title", "标题"),
    ("authors", "作者"),
    ("year", "年份"),
    ("journal", "期刊"),
    ("cited_by_count", "被引"),
    ("status", "状态"),
    ("channel", "下载通道"),
    ("file", "文件"),
    ("reason", "失败原因"),
    ("detail", "详情"),
]


def is_valid_pdf(path: Path) -> tuple[bool, str]:
    try:
        size = path.stat().st_size
    except OSError as exc:
        return False, f"无法读取：{exc}"
    if size < MIN_PDF_BYTES:
        return False, f"文件过小（{size} 字节），可能是错误页而不是 PDF"
    try:
        with path.open("rb") as fh:
            head = fh.read(5)
    except OSError as exc:
        return False, f"无法读取：{exc}"
    if head[:4] != PDF_MAGIC:
        return False, f"文件头不是 PDF（{head!r}）"
    return True, f"{size/1024:.0f} KB"


def scan_pdfs() -> dict[str, dict]:
    """把 PDF 目录扫成 {文件名主干: {path, valid, detail, size}}。"""
    C.PDF_DIR.mkdir(parents=True, exist_ok=True)
    out: dict[str, dict] = {}
    for path in sorted(C.PDF_DIR.glob("*.pdf")):
        valid, detail = is_valid_pdf(path)
        out[path.stem] = {"path": path, "valid": valid, "detail": detail,
                          "size": path.stat().st_size if path.exists() else 0}
    return out


def recover_doi_from_name(stem: str) -> str:
    """从文件名反推 DOI。实现统一在 C.recover_doi_from_filename，避免两套规则漂移。"""
    return C.recover_doi_from_filename(stem)


def consolidate_pdfs() -> dict:
    """同一 DOI 的多个副本合并。

    上游 lanes 引擎会把每条车道的中间结果也落盘，文件名带来源后缀，例如：
        10.1371_journal.pone.0173717.pdf
        10.1371_journal.pone.0173717_Sci-Hub.pdf
        10.1038_nature12373.pdf                          （CrossrefPage，259 KB）
        10.1038_nature12373_Sci-Hub_scihub_sci-hub_vg.pdf（921 KB，更完整）
    本 skill 的契约是「所有 PDF 放一个目录、用 DOI 命名」，所以这里：
      1. 同一 DOI 只保留**最完整的一份**（先看是否有效 PDF，再比大小）
      2. 把它落到规范 DOI 文件名
      3. 其余副本移入 pdf_alternates/（不删，留作对照）
    """
    C.PDF_DIR.mkdir(parents=True, exist_ok=True)
    alternates = C.ALTERNATES_DIR

    groups: dict[str, list[Path]] = {}
    for path in sorted(C.PDF_DIR.glob("*.pdf")):
        doi = C.recover_doi_from_filename(path.stem)
        if doi:
            groups.setdefault(C.normalize_doi(doi), []).append(path)

    stats = {"groups": 0, "merged": 0, "promoted": 0}

    def score(p: Path) -> tuple[int, int]:
        valid, _ = is_valid_pdf(p)
        try:
            size = p.stat().st_size
        except OSError:
            size = 0
        return (1 if valid else 0, size)

    for doi, paths in groups.items():
        if len(paths) < 2:
            continue
        stats["groups"] += 1
        canonical = C.PDF_DIR / (C.identifier_to_filename(doi) + ".pdf")
        best = max(paths, key=score)

        if best != canonical:
            if canonical.exists():
                alternates.mkdir(parents=True, exist_ok=True)
                try:
                    canonical.replace(alternates / canonical.name)
                    stats["merged"] += 1
                except OSError:
                    best = canonical      # 换不动就保持现状，不冒险
            if best != canonical:
                try:
                    best.replace(canonical)
                    stats["promoted"] += 1
                except OSError:
                    pass

        for path in paths:
            if path == canonical or not path.exists():
                continue
            alternates.mkdir(parents=True, exist_ok=True)
            target = alternates / path.name
            try:
                if target.exists():
                    path.unlink(missing_ok=True)
                else:
                    path.replace(target)
                stats["merged"] += 1
            except OSError:
                pass

    return stats


def collect(queue: dict) -> tuple[list[C.Item], list[C.Item]]:
    """返回 (本次队列的条目, 工作目录里不属于本次队列的 PDF)。

    为什么要分开：用户的工作目录是复用的。跑第二个任务、而这次清单更短时，
    上一轮下好的 PDF 会在磁盘上"无主"。早期版本把它们当成条目混进元信息表，
    结果是「我要 2 篇」的表里冒出 5 行。那些文件不该消失，但也**不属于本次交付**——
    所以只进 doi_index.json（PDF 目录的完整索引）与页脚说明，不进主表。
    """
    # 队列只代表**这一次任务**。两处过滤都必要：
    #   * origin=recovered：早期版本把上一轮捡到的"孤儿"写回队列，下次任务队列变短时
    #     它们会一直粘在表里。孤儿已在 collect() 末尾单独返回，这里从源头清掉。
    #   * selected=False：门① 没被选中的候选**不是交付物**。不滤掉的话，
    #     用户勾了 2 条、表里却出现 31 行（29 行 pending），交付物直接失真。
    items = [i for i in C.queue_items(queue, only_selected=True) if i.origin != "recovered"]
    pdfs = scan_pdfs()

    # 1) 认领磁盘文件：先按契约规范名，再按别名（上游可能用裸 arXiv 号等落盘）
    claimed: set[str] = set()
    resolved: list[C.Item] = []
    for it in items:
        entry = None
        if it.kind in ("doi", "arxiv"):
            for stem in C.alias_stems(it.id):
                if stem in pdfs:
                    entry = pdfs[stem]
                    claimed.add(stem)
                    break
        if entry is None and it.file:
            stem = Path(it.file).stem
            if stem in pdfs:
                entry = pdfs[stem]
                claimed.add(stem)

        if entry is None:
            # 上一轮收编的「孤儿」如果这一轮已被合并/清理掉，直接丢弃——
            # 那是中间产物，不是失败项。
            if it.origin == "recovered":
                continue
            if it.status == C.STATUS_OK:
                it.status = C.STATUS_FAIL
                it.reason = "file_missing"
                it.detail = "下载阶段标记成功，但磁盘上找不到对应 PDF"
                it.file = ""
            resolved.append(it)
            continue
        if entry["valid"]:
            it.status = C.STATUS_OK
            it.reason = ""
            it.file = str(entry["path"])
            it.detail = entry["detail"]
        else:
            it.status = C.STATUS_FAIL
            it.reason = "file_missing"
            it.detail = entry["detail"]
            it.file = ""
        resolved.append(it)

    items = resolved

    # 2) 工作目录里未被本次队列认领的 PDF
    orphans: list[C.Item] = []
    for stem, entry in pdfs.items():
        if stem in claimed:
            continue
        doi = recover_doi_from_name(stem)
        orphan = C.Item(
            id=doi or stem,
            kind="doi" if doi else "unknown",
            title="",
            origin="recovered",
            status=C.STATUS_OK if entry["valid"] else C.STATUS_FAIL,
            file=str(entry["path"]) if entry["valid"] else "",
            reason="" if entry["valid"] else "file_missing",
            detail="工作目录中存在，但不属于本次队列（上一轮遗留或人工放入）" if entry["valid"] else entry["detail"],
        )
        orphans.append(orphan)

    return items, orphans


def write_csv(items: list[C.Item], path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow([cn for _, cn in COLUMNS])
        for n, it in enumerate(items, 1):
            writer.writerow([
                n, it.id, it.title, "; ".join(it.authors), it.year, it.journal,
                it.cited_by_count, it.status, it.channel,
                Path(it.file).name if it.file else "",
                REASON_CN.get(it.reason, it.reason) if it.status != C.STATUS_OK else "",
                it.detail.replace("\n", " ")[:300],
            ])


def write_xlsx(items: list[C.Item], path: Path) -> bool:
    try:
        from openpyxl import Workbook  # type: ignore
        from openpyxl.styles import Font  # type: ignore
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "文献元信息"
    ws.append([cn for _, cn in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for n, it in enumerate(items, 1):
        ws.append([
            n, it.id, it.title, "; ".join(it.authors), it.year, it.journal,
            it.cited_by_count, it.status, it.channel,
            Path(it.file).name if it.file else "",
            REASON_CN.get(it.reason, it.reason) if it.status != C.STATUS_OK else "",
            it.detail.replace("\n", " ")[:300],
        ])
    widths = [5, 30, 60, 30, 7, 26, 7, 9, 22, 30, 26, 50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    wb.save(path)
    return True


def md_cell(value: str, limit: int = 0) -> str:
    """把任意文本变成**安全的单个 Markdown 表格单元格**。

    必须折叠空白：Crossref / bioRxiv 的标题经常带换行与连续空格，直接塞进表格会
    把行撑断（实测 bioRxiv 标题里的 `\\n` 让 metadata.md 断成多行，表格渲染全乱），
    而 CSV/XLSX 用标准引号转义所以看不出来。
    顺序也重要：先折叠空白、再转义竖线，否则 `\\|` 里的反斜杠会被后续处理搞乱。
    """
    text = " ".join(str(value or "").split())
    text = text.replace("|", "\\|")
    return text[:limit] if limit else text


def write_markdown(items: list[C.Item], path: Path, extra_count: int = 0) -> None:
    by = Counter(i.status for i in items)
    okk, bad, pend, skip = (by.get(C.STATUS_OK, 0), by.get(C.STATUS_FAIL, 0),
                            by.get("pending", 0), by.get(C.STATUS_SKIP, 0))
    tally = f"成功 {okk}，失败 {bad}"
    if pend:
        tally += f"，未处理 {pend}"
    if skip:
        tally += f"，跳过 {skip}"
    lines = [
        "# 文献元信息表",
        "",
        f"- 生成时间：{C.utcnow()}",
        f"- 总计：{len(items)} 篇（{tally}）",
        f"- PDF 目录：`{C.PDF_DIR}`",
        "",
        "| # | DOI | 标题 | 年份 | 期刊 | 被引 | 状态 | 通道 | 文件 |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    for n, it in enumerate(items, 1):
        lines.append(
            f"| {n} | {md_cell(it.id)} | {md_cell(it.title, 90)} | {md_cell(it.year)} | "
            f"{md_cell(it.journal, 40)} | {it.cited_by_count or ''} | {it.status} | "
            f"{md_cell(it.channel)} | {Path(it.file).name if it.file else ''} |"
        )
    if extra_count:
        lines += [
            "",
            f"> 工作目录中另有 **{extra_count}** 个 PDF 不属于本次队列（上一轮遗留或人工放入），"
            f"未列入上表；完整映射见 `doi_index.json`。",
        ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_failures(items: list[C.Item], path: Path) -> None:
    failed = [i for i in items if i.status == C.STATUS_FAIL]
    lines = [
        "# 失败清单与重试建议",
        "",
        f"- 生成时间：{C.utcnow()}",
        f"- 失败 {len(failed)} 篇 / 共 {len(items)} 篇",
        "",
    ]
    if not failed:
        lines += [
            "本次运行没有失败项。",
            "",
            "> 注意：这个文件反映的是**最近一次**运行。如果你跑过好几轮（例如先检索下载、"
            "后来又补了几篇），前面那些轮次的失败归因不会留在本文件里——它们在同目录的 "
            "`logs/*.md` 与 `failures-<时间戳>.md` 里。",
        ]
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        return

    lines += ["## 归因分布", "", "| 原因 | 含义 | 数量 | 建议 |", "|---|---|---|---|"]
    for reason, n in Counter(i.reason for i in failed).most_common():
        lines.append(f"| `{reason}` | {REASON_CN.get(reason, reason)} | {n} | {RETRY_HINT.get(reason, '')} |")

    lines += ["", "## 逐条明细", "", "| DOI | 标题 | 原因 | 原始信息 |", "|---|---|---|---|"]
    for it in failed:
        lines.append(
            f"| {md_cell(it.id)} | {md_cell(it.title, 70)} | `{it.reason}` | "
            f"{md_cell(it.detail, 220)} |"
        )

    lines += ["", "## 重试", "", "```", "python 0路由/scripts/download.py --retry-only",
              "python 0路由/scripts/finalize.py", "```"]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_doi_index(items: list[C.Item], path: Path) -> None:
    index = {}
    for it in items:
        if it.status != C.STATUS_OK or not it.file:
            continue
        index[it.id] = {
            "file": Path(it.file).name,
            "path": it.file,
            "channel": it.channel,
            "title": it.title,
            "year": it.year,
            "ts": C.utcnow(),
        }
    C.write_json(path, {
        "schema_version": C.SCHEMA_VERSION,
        "generated_at": C.utcnow(),
        "pdf_dir": str(C.PDF_DIR),
        "count": len(index),
        "entries": index,
    })


def run(queue_path: Path, out_dir: Path) -> int:
    C.ensure_dirs()
    queue = C.load_queue(queue_path)
    out_dir.mkdir(parents=True, exist_ok=True)

    C.step("P6 后处理")

    # 先合并同一 DOI 的多个副本（上游各车道的中间文件），再按磁盘实况对账
    merge_stats = consolidate_pdfs()
    if merge_stats["groups"]:
        C.info(f"同一 DOI 有多个副本的分组：{merge_stats['groups']} 组，"
               f"提升为规范文件 {merge_stats['promoted']} 个，"
               f"其余 {merge_stats['merged']} 份移入 {C.ALTERNATES_DIR}")

    items, others = collect(queue)

    if others:
        C.info(f"工作目录中另有 {len(others)} 个 PDF 不属于本次队列"
               f"（上一轮遗留或人工放入）；只进 doi_index.json，不进元信息表")

    # 文件名规范化：把任何非 DOI 命名的成功文件改名为 DOI 命名。
    # 必须覆盖 arxiv 条目——上游会把 `10.48550/arXiv.1706.03762` 归一化成裸
    # arXiv 号并落盘成 `1706.03762.pdf`，直接违反本 skill 的 DOI 命名契约。
    renamed = 0
    for it in items:
        if it.status != C.STATUS_OK or not it.file or not it.id:
            continue
        if it.kind not in ("doi", "arxiv"):
            continue
        want = C.PDF_DIR / (C.identifier_to_filename(it.id) + ".pdf")
        cur = Path(it.file)
        if cur == want or not cur.exists():
            it.file = str(want if want.exists() else cur)
            continue
        if want.exists():
            # 规范名已存在（同一篇重复下载）：保留规范名的那份，删掉别名副本
            try:
                if cur.stat().st_size == want.stat().st_size:
                    cur.unlink(missing_ok=True)
                    renamed += 1
            except OSError:
                pass
            it.file = str(want)
            continue
        try:
            cur.replace(want)
            it.file = str(want)
            renamed += 1
        except OSError:
            pass
    if renamed:
        C.info(f"规范化文件名为 DOI 命名：{renamed} 个")

    valid = [i for i in items if i.status == C.STATUS_OK]
    failed = [i for i in items if i.status == C.STATUS_FAIL]
    skipped = [i for i in items if i.status == C.STATUS_SKIP]

    C.step("P7 产出")
    csv_path = out_dir / "metadata.csv"
    xlsx_path = out_dir / "metadata.xlsx"
    md_path = out_dir / "metadata.md"
    fail_path = out_dir / "failures.md"
    index_path = out_dir / "doi_index.json"

    write_csv(items, csv_path); C.ok(f"元信息表（CSV）：{csv_path}")
    if write_xlsx(items, xlsx_path):
        C.ok(f"元信息表（Excel）：{xlsx_path}")
    else:
        C.warn("没有 openpyxl，跳过 .xlsx（用 pipeline.py 跑会自动切到 venv 解释器）")
    write_markdown(items, md_path, extra_count=len(others)); C.ok(f"元信息表（Markdown）：{md_path}")
    write_failures(items, fail_path); C.ok(f"失败清单与建议：{fail_path}")

    # failures.md 反映的是"最近一次运行"。多轮跑（先下载、后补几篇）时后一轮会把它覆盖，
    # 前几轮的失败归因就没了。只要这轮有失败项，就额外留一份带时间戳的快照做审计线索。
    if failed:
        snap = out_dir / f"failures-{C.run_stamp()}.md"
        write_failures(items, snap)
        C.ok(f"本轮失败快照（不会被后续轮次覆盖）：{snap.name}")

    write_doi_index(items + others, index_path); C.ok(f"DOI → 文件索引：{index_path}")

    C.save_items(queue, items, queue_path)

    summary = (f"总计 {len(items)}：成功 {len(valid)}，失败 {len(failed)}"
               + (f"，跳过 {len(skipped)}" if skipped else ""))
    log = C.RunLog("finalize")
    for it in items:
        log.item(it.id, it.status, it.reason, it.detail, channel=it.channel, file=it.file)
    log_path = log.close(summary=summary)
    C.ok(f"运行日志：{log_path}")

    C.step("结果")
    C.ok(summary)
    if failed:
        C.bad("失败归因：")
        for reason, n in Counter(i.reason for i in failed).most_common():
            C.info(f"  {reason:<22} {n} 条 — {REASON_CN.get(reason, '')}")
        C.info("重试：python 0路由/scripts/download.py --retry-only")
    return 0 if not failed else 1


def main() -> int:
    C.console()
    ap = argparse.ArgumentParser(description="P6/P7 · 后处理与产出")
    ap.add_argument("--queue", default=str(C.DEFAULT_QUEUE))
    ap.add_argument("--out-dir", default="", help="交付物目录（默认 skill 的 workspace/pdfs）")
    args = ap.parse_args()
    if args.out_dir:
        C.set_output_dir(args.out_dir)
    return run(Path(args.queue), Path(args.out_dir) if args.out_dir else C.PDF_DIR)


if __name__ == "__main__":
    raise SystemExit(main())
